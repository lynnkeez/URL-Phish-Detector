from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "PhishGuard_Testing_and_Evaluation_Report.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
PALE_BLUE = "E8EEF5"
PALE_GRAY = "F2F4F7"
INK = RGBColor(11, 37, 69)


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_width(cell, width):
    properties = cell._tc.get_or_add_tcPr()
    width_element = properties.first_child_found_in("w:tcW")
    if width_element is None:
        width_element = OxmlElement("w:tcW")
        properties.append(width_element)
    width_element.set(qn("w:w"), str(int(width * 1440)))
    width_element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table_properties = table._tbl.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), "9360")
    table_width.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    table_properties.append(indent)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)


def style_run(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(doc, text="", bold_prefix=None, italic=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        style_run(paragraph.add_run(bold_prefix), bold=True)
        style_run(paragraph.add_run(text[len(bold_prefix):]), italic=italic)
    else:
        style_run(paragraph.add_run(text), italic=italic)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.167
    style_run(paragraph.add_run(text))
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    style_run(run, size={1: 16, 2: 13, 3: 12}[level], bold=True, color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level])
    return paragraph


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0].cells
    for cell, value in zip(header, headers):
        set_cell_shading(cell, PALE_BLUE)
        paragraph = cell.paragraphs[0]
        style_run(paragraph.add_run(value), size=10, bold=True, color=DARK_BLUE)
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            paragraph = cell.paragraphs[0]
            style_run(paragraph.add_run(value), size=9.5)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_GRAY)
    paragraph = cell.paragraphs[0]
    style_run(paragraph.add_run(title + " "), bold=True, color=DARK_BLUE)
    style_run(paragraph.add_run(text))
    set_table_geometry(table, [6.5])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_reference(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.10
    style_run(paragraph.add_run(text))


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for level, size, before, after, color in [(1, 16, 16, 8, BLUE), (2, 13, 12, 6, BLUE), (3, 12, 8, 4, DARK_BLUE)]:
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_run(header.add_run("PHISHGUARD | TESTING AND EVALUATION REPORT"), size=9, bold=True, color="666666")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(footer.add_run("APT3065 Coursework | Summer 2026"), size=9, color="666666")


def build():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(16)
    title.paragraph_format.space_after = Pt(4)
    style_run(title.add_run("PHISHGUARD"), size=23, bold=True, color="000000")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    style_run(subtitle.add_run("Testing and Evaluation Report"), size=14, color="555555")
    for label, value in [("Module", "APT3065"), ("Project", "URL Phishing Guard"), ("Assessment focus", "Testing strategy, execution evidence, and technical evaluation"), ("Prepared by", "[Insert student name and registration number]")]:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        style_run(paragraph.add_run(label + ": "), bold=True)
        style_run(paragraph.add_run(value))
    doc.add_paragraph()
    add_callout(doc, "Purpose.", "This report defines the verification approach for PhishGuard and records the evidence required to demonstrate reliable, robust behaviour. It should be submitted with the generated test logs, performance evidence, screenshots, and final model metrics.")

    add_heading(doc, "1. Testing Strategy and Planning")
    add_paragraph(doc, "PhishGuard is a Flask application that estimates phishing risk using URL-only lexical and structural features. The strategy separates verification into unit, integration, system, and acceptance levels so that defects can be isolated and explained, consistent with the general concepts of structured software testing (ISO/IEC/IEEE, 2022). The scope covers feature extraction, model integration, web/API behaviour, session history, and response time. It excludes website crawling because the design intentionally performs no external network requests.")
    add_table(doc, ["Level", "Objective", "Method and tool", "Pass criterion"], [
        ("Unit", "Verify individual validation and feature rules.", "Python unittest; direct function calls.", "Every assertion passes."),
        ("Integration", "Verify Flask routes, JSON API, redirects, and session state.", "Flask test client with mocked inference.", "Correct status code, response data, and session change."),
        ("System", "Verify feature extraction and a real trained model work together.", "unittest against the trained model.", "Complete result schema; each inference below 2 seconds."),
        ("Acceptance", "Verify end-user journeys in a desktop browser.", "Manual checklist and screenshots.", "All critical user scenarios accepted."),
    ], [1.0, 1.65, 2.15, 1.7])

    add_heading(doc, "1.1 Test objectives, scope, and criteria", level=2)
    add_bullet(doc, "Correctness: valid URLs are normalised, invalid input is rejected safely, and the output contains a label, confidence, risk level, and explanation.")
    add_bullet(doc, "Robustness: blank, unsupported-scheme, whitespace-containing, malformed, boundary-length, and long historical training URLs are handled without application failure.")
    add_bullet(doc, "Performance: prediction processing is measured against a maximum response-time requirement of 2,000 ms per URL after model loading.")
    add_bullet(doc, "Traceability: every test is linked to an explicit requirement and produces repeatable evidence rather than an informal claim.")

    add_heading(doc, "2. Test Cases and Coverage")
    add_table(doc, ["ID", "Level", "Input / scenario", "Expected result"], [
        ("UT-01", "Unit", "example.com/login", "Scheme is added and the 31-feature order is stable."),
        ("UT-02", "Unit", "paypal.login.verify.example.zip/%41", "Suspicious TLD, brand-in-subdomain, keyword, and encoding indicators are detected."),
        ("UT-03", "Unit", "Blank, FTP, or whitespace URL", "ValueError returned; no crash."),
        ("UT-04", "Unit", "255- and 256-character web inputs", "255 characters accepted; 256 rejected."),
        ("UT-05", "Unit", "Historical URL longer than 255 characters", "Training feature extraction accepts it."),
        ("IT-01", "Integration", "POST /api/check with valid JSON", "HTTP 200 and prediction JSON returned."),
        ("IT-02", "Integration", "POST /api/check with missing URL", "HTTP 400 and a clear error message."),
        ("IT-03", "Integration", "Submit form then inspect history", "Redirect occurs and one session-history entry exists."),
        ("IT-04", "Integration", "Clear history", "HTTP 204 and session history removed."),
        ("ST-01", "System", "Safe-looking, suspicious, and IP-host URLs", "Valid result schema; confidence from 0 to 100; each result under 2 s."),
    ], [0.7, 0.8, 2.45, 2.55])

    add_heading(doc, "3. Test Execution and Evidence")
    add_paragraph(doc, "The project stores executable tests in the tests directory. Unit tests are intentionally independent of a trained model. Integration tests mock inference so a web-route defect is not hidden by model availability. The API and input-validation cases support a methodical assessment of web-application controls and error handling (OWASP Foundation, 2020; Scarfone et al., 2008). System tests use the real model and are skipped, not passed, until retraining produces a compatible model.")
    add_callout(doc, "Current verified result.", "The URL feature test suite executed successfully: 5 tests run, 5 passed, 0 failed. These tests cover normalisation, phishing indicators, invalid input, web-input boundary behaviour, and long historical training URLs.")
    add_paragraph(doc, "The following commands generate the evidence to attach to this report after dependencies are installed and the model is retrained:")
    code = doc.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.25)
    code.paragraph_format.space_after = Pt(8)
    style_run(code.add_run("python -m unittest discover -s tests -v > test_evidence\\unit-integration-system.log 2>&1\ntype test_evidence\\unit-integration-system.log\npython scripts\\run_performance_check.py"), size=9.5, color="1F4D78")
    add_table(doc, ["Evidence item", "Location", "Purpose"], [
        ("Automated test log", "test_evidence/unit-integration-system.log", "Proves test names, pass/fail/skip state, and execution date."),
        ("Performance output", "test_evidence/performance.json", "Records individual cases, mean/max latency, threshold, and pass state."),
        ("Acceptance screenshots", "test_evidence/AT-01 to AT-05", "Shows user-facing behaviour in a browser."),
        ("Model metrics", "models/metrics.json", "Provides held-out accuracy, precision, recall, F1, ROC-AUC, and confusion matrix."),
    ], [1.55, 2.35, 2.6])

    add_heading(doc, "4. Results Analysis and Technical Evaluation")
    add_paragraph(doc, "The initial implementation revealed two defects during verification. First, the original model had been evaluated using webpage-content features that the live URL-only application could not calculate. This would have overstated real-world performance. The corrected design uses one shared URL-only feature pipeline for training, testing, and inference. Maintaining this alignment reduces a form of machine-learning technical debt caused by inconsistent data dependencies between development and deployment (Sculley et al., 2015). Second, a 255-character limit introduced for browser-session safety was incorrectly applied to the historical training set. The training extractor now accepts long historical records while retaining the web-input limit.")
    add_table(doc, ["Finding", "Underlying cause", "Corrective action", "Verification"], [
        ("Unreliable 100% model score", "Training used content features unavailable in live inference.", "Retrained design uses only 31 URL-derived features.", "Compare held-out metrics after retraining."),
        ("Training stopped on long URL", "Web-input boundary rule was applied to dataset rows.", "Training calls extraction with length enforcement disabled and skips malformed records.", "UT-05 passed."),
        ("Python 3.14 installation failure", "Old pandas pin lacked a compatible Windows wheel and pip attempted a source build.", "Requirements updated to Python 3.14-compatible releases.", "Reinstall and retain pip log."),
    ], [1.45, 1.75, 2.0, 1.3])
    add_heading(doc, "4.1 Quantitative evaluation to complete after retraining", level=2)
    add_table(doc, ["Measure", "Requirement / interpretation", "Final measured value"], [
        ("Accuracy", "Overall held-out classification correctness.", "[Insert from metrics.json]"),
        ("Precision", "Of URLs classified phishing, the proportion truly phishing.", "[Insert from metrics.json]"),
        ("Recall", "Of phishing URLs, the proportion detected. Low recall increases security risk.", "[Insert from metrics.json]"),
        ("F1 score", "Balance of precision and recall for the phishing class.", "[Insert from metrics.json]"),
        ("ROC-AUC", "Probability ranking quality across decision thresholds.", "[Insert from metrics.json]"),
        ("Maximum latency", "Must be below 2,000 ms per prediction.", "[Insert from performance.json]"),
    ], [1.1, 3.5, 1.95])

    add_heading(doc, "5. Risks, Limitations, and Acceptance")
    add_bullet(doc, "False negatives are the principal security risk because a phishing URL could be shown as legitimate. False positives are a usability risk because they may block harmless links.")
    add_bullet(doc, "A URL-only classifier cannot reliably identify every compromised legitimate domain. HTTPS encrypts traffic but does not prove that a website is trustworthy.")
    add_bullet(doc, "The domain-grouped split is deliberately stricter than a random split. It reduces domain leakage and produces a more credible estimate of generalisation, even if the final score is lower.")
    add_bullet(doc, "Model performance can drift as attacker patterns change; retraining and a dated metrics file are therefore required before final submission (Sculley et al., 2015).")
    add_heading(doc, "5.1 Acceptance checklist", level=2)
    add_table(doc, ["ID", "User scenario", "Acceptance criterion", "Result / evidence"], [
        ("AT-01", "Valid URL", "Verdict, confidence, and indicators shown.", "[Screenshot / Pass-Fail]"),
        ("AT-02", "URL without scheme", "Normalised and classified.", "[Screenshot / Pass-Fail]"),
        ("AT-03", "Invalid, blank, FTP, or overlong input", "Clear message; no crash.", "[Screenshot / Pass-Fail]"),
        ("AT-04", "View and clear session history", "History is shown and clear removes it.", "[Screenshot / Pass-Fail]"),
        ("AT-05", "Browser navigation", "Check, result, history, and About pages usable.", "[Screenshot / Pass-Fail]"),
        ("AT-06", "POST /api/check", "Valid JSON result; missing URL gives HTTP 400.", "[Log / Pass-Fail]"),
    ], [0.7, 1.8, 2.35, 1.7])
    add_callout(doc, "Submission checklist.", "Before submitting, retrain the model, insert the generated metrics and performance values above, attach the test log and screenshots, complete the acceptance results, and explain any failed or skipped tests honestly.")

    doc.add_page_break()
    add_heading(doc, "References")
    add_reference(doc, "International Organization for Standardization, International Electrotechnical Commission, & Institute of Electrical and Electronics Engineers. (2022). ISO/IEC/IEEE 29119-1:2022 software and systems engineering - Software testing - Part 1: General concepts. https://www.iso.org/standard/81291.html")
    add_reference(doc, "OWASP Foundation. (2020). OWASP web security testing guide (Version 4.2). https://owasp.org/www-project-web-security-testing-guide/")
    add_reference(doc, "Scarfone, K., Souppaya, M., Cody, A., & Orebaugh, A. (2008). Technical guide to information security testing and assessment (NIST Special Publication 800-115). National Institute of Standards and Technology. https://doi.org/10.6028/NIST.SP.800-115")
    add_reference(doc, "Sculley, D., Holt, G., Golovin, D., Davydov, E., Phillips, T., Ebner, D., Chaudhary, V., Young, M., Crespo, J.-F., & Dennison, D. (2015). Hidden technical debt in machine learning systems. In C. Cortes et al. (Eds.), Advances in neural information processing systems (Vol. 28). https://papers.nips.cc/paper/5656-hidden-technical-debt-in-machine-learning-systems.pdf")

    doc.core_properties.title = "PhishGuard Testing and Evaluation Report"
    doc.core_properties.subject = "APT3065 coursework testing evidence"
    doc.core_properties.author = "PhishGuard project team"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
